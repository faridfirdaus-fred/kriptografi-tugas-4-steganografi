from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
ASSETS = OUT / "report_assets"
REPORT = ROOT / "237006081_Farid Firdaus_Praktikum.docx"
RESULTS = json.loads((OUT / "hasil_eksperimen.json").read_text(encoding="utf-8"))


def set_font(run, name="Times New Roman", size=12, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_field(paragraph):
    paragraph.add_run("Halaman ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def resize_preview(source: Path, name: str, max_width=1400, max_height=1000) -> Path:
    ASSETS.mkdir(exist_ok=True)
    target = ASSETS / name
    with Image.open(source) as image:
        image = image.convert("RGB")
        image.thumbnail((max_width, max_height), Image.Resampling.LANCZOS)
        image.save(target, quality=90, optimize=True)
    return target


def side_by_side(left: Path, right: Path, name: str, left_label: str, right_label: str) -> Path:
    ASSETS.mkdir(exist_ok=True)
    target = ASSETS / name
    with Image.open(left) as a, Image.open(right) as b:
        a, b = a.convert("RGB"), b.convert("RGB")
        a.thumbnail((900, 700), Image.Resampling.LANCZOS)
        b.thumbnail((900, 700), Image.Resampling.LANCZOS)
        w, h = a.width + b.width + 60, max(a.height, b.height) + 70
        canvas = Image.new("RGB", (w, h), "white")
        canvas.paste(a, (20, 45))
        canvas.paste(b, (a.width + 40, 45))
        draw = ImageDraw.Draw(canvas)
        draw.text((20, 15), left_label, fill="black")
        draw.text((a.width + 40, 15), right_label, fill="black")
        canvas.save(target, quality=90, optimize=True)
    return target


def shade(cell, fill="D9EAF7"):
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    properties.append(element)


def set_cell(cell, value, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(value))
    set_font(run, size=10, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, item in enumerate(headers):
        set_cell(t.rows[0].cells[i], item, True)
        shade(t.rows[0].cells[i])
    for row in rows:
        cells = t.add_row().cells
        for i, item in enumerate(row):
            set_cell(cells[i], item)
    doc.add_paragraph()
    return t


def paragraph(doc, text="", bold_prefix=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=3):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(after)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def title(doc, text):
    p = paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.LEFT, after=4)
    for r in p.runs:
        r.bold = True
    return p


def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(.5)
    for line in text.strip("\n").splitlines():
        r = p.add_run(line + "\n")
        set_font(r, name="Courier New", size=8.5)
    return p


def figure(doc, image: Path, caption: str, width=15.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image), width=Cm(width))
    c = paragraph(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER, after=5)
    for r in c.runs:
        r.italic = True
        r.font.size = Pt(10)


def main():
    # Only compact previews are embedded. Original experiment artifacts remain in output/.
    cover_preview = resize_preview(ROOT / "cover_farid.png", "cover_preview.jpg")
    p1_compare = side_by_side(ROOT / "cover_farid.png", OUT / "p1" / "stego_p1_fixed_header32.png", "p1_cover_stego.jpg", "Cover", "Stego P1 (perbaikan)")
    map_compare = side_by_side(OUT / "p2" / "change_map_sequential.png", OUT / "p2" / "change_map_random.png", "p2_maps.jpg", "Peta perubahan sequential", "Peta perubahan acak")
    lsb_compare = side_by_side(OUT / "bonus" / "lsb_plane_cover.png", OUT / "bonus" / "lsb_plane_stego_50_percent.png", "bonus_lsb_planes.jpg", "Bidang LSB cover", "Bidang LSB stego 50%")
    charts = resize_preview(OUT / "p3" / "grafik_psnr.png", "p3_grafik_psnr.jpg")
    montage = resize_preview(OUT / "p3" / "perbandingan_stego_m1_m4.png", "p3_montage.jpg")

    doc = Document()
    section = doc.sections[0]
    for margin in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, margin, Cm(2))
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(3)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(footer)

    p = paragraph(doc, "TUGAS 4 - PRAKTIKUM STEGANOGRAFI", align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    p.runs[0].bold = True
    p = paragraph(doc, "BAGIAN II - TUGAS PRAKTIKUM", align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    p.runs[0].bold = True
    paragraph(doc, "Nama: Farid Firdaus", align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    paragraph(doc, "NIM: 237006081", align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

    title(doc, "A. Lingkungan dan Bahan Praktikum")
    paragraph(doc, "Praktikum menggunakan Python 3 dengan NumPy, Pillow, dan Matplotlib. Citra cover adalah foto milik sendiri berformat PNG RGB berukuran 4032 × 3024 piksel. Pesan rahasia yang digunakan memenuhi format NIM - nama lengkap - kalimat bebas minimal 10 kata.")
    table(doc, ["Komponen", "Nilai"], [
        ["Pesan", RESULTS["message"]],
        ["Seed posisi acak", "237006081"],
        ["Seed salah untuk pengujian", "237006082"],
        ["Seed enkripsi PRNG", "173006081"],
        ["Kapasitas teoretis 1-bit LSB", f"{RESULTS['cover']['capacity_bits_1_lsb']:,} bit".replace(",", ".")],
    ])
    figure(doc, cover_preview, "Gambar 1. Citra cover milik sendiri yang digunakan pada praktikum.", 14)

    title(doc, "P1. Analisis dan Perbaikan Program Slide")
    paragraph(doc, "Program slide direproduksi terlebih dahulu. Dua masalah ditemukan. Pertama, kode memakai bin(x)[2:9] tanpa padding 8-bit. Pada nilai kanal rendah, potongan bit tersebut menggeser posisi LSB sehingga perubahan dapat lebih besar dari satu. Kedua, range(0, 2) hanya menggunakan kanal R dan G, bukan tiga kanal RGB.")
    code(doc, "# Perbaikan penyisipan 1-bit LSB pada semua kanal RGB\nflat[:len(bits)] = (flat[:len(bits)] & 0xFE) | bits\n# Header: 32 bit panjang payload, menggantikan delimiter \"stego\"")
    p1 = RESULTS["p1"]
    table(doc, ["Pengujian", "Selisih maksimum", "Channel-byte berubah", "Keterangan"], [
        ["Baseline slide pada foto", p1["baseline_max_difference"], p1["baseline_changed_bytes"], "Bug tersamarkan oleh kanal awal yang terang"],
        ["Baseline slide data sintetis rendah", p1["synthetic_low_value"]["baseline_max_difference"], p1["synthetic_low_value"]["baseline_changed_bytes"], "Bug bin(x)[2:9] terbukti muncul"],
        ["Perbaikan pada data sintetis", p1["synthetic_low_value"]["fixed_max_difference"], p1["synthetic_low_value"]["fixed_changed_bytes"], "Sesuai teori 1-bit LSB"],
        ["Perbaikan pada foto", p1["fixed_max_difference"], p1["fixed_changed_bytes"], "Sesuai teori 1-bit LSB"],
    ])
    paragraph(doc, "Selisih pada tabel dihitung per channel-byte dengan np.abs(cover.astype(int) - stego.astype(int)). Teori 1-bit LSB menyatakan perubahan maksimum adalah satu. Kapasitas program slide adalah R dan G saja, yaitu 24.385.536 bit. Setelah memakai R, G, dan B, kapasitasnya menjadi 36.578.304 bit.")
    paragraph(doc, "Delimiter \"stego\" tidak aman karena kata tersebut mungkin muncul sebagai bagian dari pesan. Perbaikan menggunakan header 32-bit yang memuat panjang payload, sehingga ekstraksi membaca jumlah bit yang pasti.")
    figure(doc, p1_compare, "Gambar 2. Perbandingan cover dan stego P1 setelah perbaikan.", 15.5)
    figure(doc, ROOT / "screenshots" / "p1_console.png", "Gambar 3. Output P1: baseline, bukti sintetis, kapasitas, dan hasil perbaikan.", 15.5)

    doc.add_page_break()
    title(doc, "P2. LSB Acak dengan Stego-key")
    paragraph(doc, "Metode acak memakai np.random.default_rng(seed).permutation(...) untuk membentuk urutan indeks channel-byte tanpa duplikasi. Seed 237006081 menjadi stego-key posisi. Header 32-bit dan payload ditempatkan pada indeks awal urutan permutasi tersebut.")
    table(doc, ["Skenario", "Hasil"], [
        ["Ekstraksi dengan seed 237006081", "Berhasil, pesan asli pulih"],
        ["Ekstraksi dengan seed 237006082", RESULTS["p2"]["wrong_seed_status"]],
        ["Channel-byte berubah sequential", RESULTS["p2"]["sequential_changed_bytes"]],
        ["Channel-byte berubah acak", RESULTS["p2"]["random_changed_bytes"]],
    ])
    paragraph(doc, "Jumlah perubahan hampir sama karena payload identik. Perbedaannya ada pada lokasi: sequential terkonsentrasi pada awal raster, sedangkan acak menyebar di seluruh citra. Pola acak lebih sulit dideteksi oleh inspeksi visual sederhana karena tidak membentuk area perubahan terkonsentrasi.")
    figure(doc, map_compare, "Gambar 4. Peta perubahan putih: sequential terkonsentrasi, acak tersebar.", 15.5)
    figure(doc, ROOT / "screenshots" / "p2_console.png", "Gambar 5. Output ekstraksi seed benar dan seed salah pada P2.", 15.5)

    title(doc, "P3. Eksperimen m-bit LSB dan Kualitas Citra")
    paragraph(doc, "Setiap nilai m menggunakan payload acak sebesar 100% kapasitas. Header sengaja tidak dipakai pada P3 karena seluruh kapasitas dialokasikan untuk eksperimen MSE dan PSNR; citra P3 bukan artefak round-trip payload. MSE dihitung dengan mean((cover - stego)^2), sedangkan PSNR = 10 log10(255²/MSE).")
    table(doc, ["m", "Kapasitas (KB)", "MSE", "PSNR (dB)"], [[x["m"], f"{x['capacity_kb']:.3f}", f"{x['mse']:.6f}", f"{x['psnr_db']:.4f}"] for x in RESULTS["p3"]])
    paragraph(doc, "Pada m=1 dan m=2, PSNR berada di atas 40 dB sehingga distorsi secara umum tidak terlihat. Pada m=3 PSNR turun menjadi 37,89 dB dan pada m=4 menjadi 31,81 dB. Dengan payload 100% kapasitas, distorsi mulai lebih mungkin terlihat pada m=3 dan semakin jelas pada m=4. Hasil ini menunjukkan trade-off: semakin besar m, kapasitas meningkat linear, tetapi fidelity menurun.")
    figure(doc, charts, "Gambar 6. Grafik PSNR terhadap nilai m.", 13)
    figure(doc, montage, "Gambar 7. Perbandingan stego-image m=1 sampai m=4 pada payload 100% kapasitas.", 15.5)
    figure(doc, ROOT / "screenshots" / "p3_console.png", "Gambar 8. Output metrik MSE dan PSNR P3.", 15.5)

    doc.add_page_break()
    title(doc, "P4. Kombinasi Kriptografi dan Steganografi")
    paragraph(doc, "Payload terlebih dahulu dienkripsi menggunakan XOR dengan keystream PRNG. Dua key digunakan: seed posisi 237006081 untuk urutan posisi LSB, dan seed enkripsi 173006081 untuk keystream. Ciphertext kemudian disisipkan memakai LSB acak.")
    table(doc, ["Skenario", "Hasil"], [
        ["Key posisi benar + key enkripsi benar", "Pesan asli berhasil dipulihkan"],
        ["Key posisi benar + key enkripsi salah", "Ciphertext terbaca, tetapi dekripsi menghasilkan byte acak"],
        ["Key posisi salah", RESULTS["p4"]["wrong_position_status"]],
    ])
    paragraph(doc, "Skenario pertama membuktikan dua key yang benar diperlukan untuk mengambil plaintext. Pada skenario kedua, posisi benar hanya menghasilkan ciphertext; tanpa key enkripsi, isi tetap tidak bermakna. Pada skenario ketiga, urutan indeks salah sehingga header panjang pun tidak valid.")
    figure(doc, ROOT / "screenshots" / "p4_console.png", "Gambar 9. Output tiga skenario key pada P4.", 15.5)

    title(doc, "P5. Pengaruh Format File")
    paragraph(doc, "Stego-image yang sama disimpan sebagai PNG, BMP, dan JPEG quality 95. Bit error rate (BER) dihitung dari persentase bit payload hasil ekstraksi yang salah dibandingkan payload asli.")
    table(doc, ["Format", "BER", "Status ekstraksi"], [[x["format"], f"{x['ber_percent']:.6f}%", x["header_status"]] for x in RESULTS["p5"]])
    paragraph(doc, "PNG dan BMP mempertahankan nilai byte piksel sehingga BER-nya nol dan pesan dapat diekstrak. JPEG memakai kompresi lossy berbasis transformasi DCT serta kuantisasi. Proses tersebut mengubah nilai kanal pada ranah spasial, termasuk LSB yang membawa bit pesan. Karena itu JPEG menghasilkan BER 48,135965% dan header panjang tidak lagi valid.")
    figure(doc, ROOT / "screenshots" / "p5_console.png", "Gambar 10. Output BER dan ekstraksi P5.", 15.5)

    title(doc, "Bonus. Steganalisis Visual")
    paragraph(doc, "Bidang LSB dibuat dengan operasi (citra & 1) × 255 untuk cover dan stego sequential berpayload 50% kapasitas. Pesan teks tidak tampak secara langsung pada bidang LSB karena bit payload membentuk pola biner, bukan bentuk huruf visual. Namun, pengisian LSB sequential dalam jumlah besar dapat mengubah distribusi bit dan menjadi sinyal bagi steganalisis statistik seperti uji chi-square. Karena itu, posisi acak pada P2 membantu menyebarkan pola perubahan.")
    figure(doc, lsb_compare, "Gambar 11. Bidang bit LSB cover dan stego sequential 50% kapasitas.", 15.5)

    title(doc, "Kesimpulan")
    paragraph(doc, "LSB 1-bit dapat menyisipkan pesan dengan perubahan channel-byte maksimum satu dan dapat diekstrak kembali secara benar pada PNG maupun BMP. Penyisipan acak menambah stego-key posisi dan menyebarkan perubahan. Penambahan jumlah bit LSB meningkatkan kapasitas, tetapi menurunkan PSNR. Enkripsi sebelum penyisipan melindungi isi pesan bila lokasi sisipan ditemukan. JPEG tidak sesuai untuk LSB spasial karena kompresi lossy merusak bit rendah.")

    doc.save(REPORT)
    print(REPORT)
    print("paragraphs=", len(doc.paragraphs), "tables=", len(doc.tables), "assets=", len(list(ASSETS.iterdir())))


if __name__ == "__main__":
    main()
